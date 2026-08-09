"""Unit tests for cover letter rendering — offline, no Playwright."""

from __future__ import annotations

from jobhunter.adapters.rendering.cover_letter_renderer import (
    to_docx,
    to_html,
    to_markdown,
)
from jobhunter.domain.cover_letter import CoverLetter


def _sample_letter() -> CoverLetter:
    return CoverLetter(
        salutation="Dear Hiring Manager,",
        opening="I am writing to apply for the Backend Developer role at TechCo.",
        body_paragraphs=[
            "In my current role, I have built scalable Python services.",
            "I am eager to bring this experience to your team.",
        ],
        closing="Thank you for your consideration. I look forward to hearing from you.",
        signature="Sincerely, Jane Doe",
        company_name="TechCo",
        role="Backend Developer",
    )


class TestToMarkdown:
    def test_contains_salutation(self) -> None:
        md = to_markdown(_sample_letter())
        assert "Dear Hiring Manager," in md

    def test_contains_body_paragraphs(self) -> None:
        md = to_markdown(_sample_letter())
        assert "scalable Python services" in md
        assert "bring this experience" in md

    def test_contains_signature(self) -> None:
        md = to_markdown(_sample_letter())
        assert "Sincerely, Jane Doe" in md

    def test_paragraph_separation(self) -> None:
        md = to_markdown(_sample_letter())
        assert "\n\n" in md


class TestToDocx:
    def test_produces_bytes(self) -> None:
        data = to_docx(_sample_letter())
        assert isinstance(data, bytes)
        assert len(data) > 100

    def test_valid_docx_header(self) -> None:
        data = to_docx(_sample_letter())
        assert data[:2] == b"PK"

    def test_no_tables_in_docx(self) -> None:
        from io import BytesIO

        from docx import Document

        data = to_docx(_sample_letter())
        doc = Document(BytesIO(data))
        assert len(doc.tables) == 0

    def test_text_present_in_docx(self) -> None:
        from io import BytesIO

        from docx import Document

        data = to_docx(_sample_letter())
        doc = Document(BytesIO(data))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Dear Hiring Manager," in full_text
        assert "scalable Python services" in full_text
        assert "Sincerely, Jane Doe" in full_text


class TestToHtml:
    def test_valid_html_structure(self) -> None:
        html = to_html(_sample_letter())
        assert "<!DOCTYPE html>" in html
        assert "</body></html>" in html

    def test_contains_letter_text(self) -> None:
        html = to_html(_sample_letter())
        assert "Dear Hiring Manager," in html
        assert "scalable Python services" in html

    def test_escapes_special_chars(self) -> None:
        letter = CoverLetter(
            salutation="Dear <Team>,",
            opening="Testing & escaping",
            body_paragraphs=["Para with <script>"],
            closing="End.",
            signature="Signed",
            company_name="A&B",
            role="Dev",
        )
        html = to_html(letter)
        assert "&lt;Team&gt;" in html
        assert "&amp;" in html
        assert "&lt;script&gt;" in html
