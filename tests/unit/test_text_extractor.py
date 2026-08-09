"""Unit tests for the text extractor adapter — DOCX and PDF extraction."""

import io

import pytest
from docx import Document
from pypdf import PdfWriter

from jobhunter.adapters.parsing.text_extractor import extract_text


def _make_docx(text: str) -> bytes:
    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_pdf(text: str) -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    # pypdf doesn't easily write text to pages, so we use reportlab-free approach:
    # create a PDF with metadata only. For a real test we embed text via annotation.
    # Instead, use a minimal valid PDF with the text embedded directly.
    from pypdf._page import PageObject

    page = PageObject.create_blank_page(width=612, height=792)
    # Direct content stream with text
    from pypdf.generic import (
        DecodedStreamObject,
        DictionaryObject,
        NameObject,
    )

    font_dict = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    resources = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_dict}),
        }
    )

    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
    stream = DecodedStreamObject()
    stream.set_data(content.encode())

    page[NameObject("/Resources")] = resources
    page[NameObject("/Contents")] = stream

    writer2 = PdfWriter()
    writer2.add_page(page)
    buf = io.BytesIO()
    writer2.write(buf)
    return buf.getvalue()


class TestDocxExtraction:
    def test_extracts_paragraphs(self) -> None:
        content = "Jane Doe\nSenior Engineer\nPython, SQL"
        docx_bytes = _make_docx(content)
        result = extract_text(docx_bytes, "resume.docx")
        assert "Jane Doe" in result
        assert "Senior Engineer" in result
        assert "Python, SQL" in result

    def test_empty_paragraphs_skipped(self) -> None:
        doc = Document()
        doc.add_paragraph("Hello")
        doc.add_paragraph("")
        doc.add_paragraph("World")
        buf = io.BytesIO()
        doc.save(buf)
        result = extract_text(buf.getvalue(), "test.docx")
        assert "Hello" in result
        assert "World" in result


class TestPdfExtraction:
    def test_extracts_text(self) -> None:
        pdf_bytes = _make_pdf("John Smith Software Engineer")
        result = extract_text(pdf_bytes, "cv.pdf")
        assert "John Smith" in result or "Software Engineer" in result

    def test_blank_pdf(self) -> None:
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        buf = io.BytesIO()
        writer.write(buf)
        result = extract_text(buf.getvalue(), "blank.pdf")
        assert result == ""


class TestUnsupportedFormat:
    def test_raises_for_txt(self) -> None:
        with pytest.raises(ValueError, match="Unsupported"):
            extract_text(b"hello", "notes.txt")

    def test_raises_for_png(self) -> None:
        with pytest.raises(ValueError, match="Unsupported"):
            extract_text(b"\x89PNG", "image.png")
