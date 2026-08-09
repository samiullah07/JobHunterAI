"""Résumé rendering — JSON, Markdown, DOCX, PDF outputs."""

from __future__ import annotations

from jobhunter.domain.resume import TailoredResume


def to_json(tailored: TailoredResume) -> str:
    return tailored.model_dump_json(indent=2)


def to_markdown(tailored: TailoredResume) -> str:
    lines: list[str] = []
    lines.append(f"# {tailored.full_name}")
    contact_parts: list[str] = []
    if tailored.email:
        contact_parts.append(tailored.email)
    if tailored.phone:
        contact_parts.append(tailored.phone)
    if tailored.location:
        contact_parts.append(tailored.location)
    if contact_parts:
        lines.append(" | ".join(contact_parts))
    if tailored.linkedin_url:
        lines.append(f"LinkedIn: {tailored.linkedin_url}")
    if tailored.github_url:
        lines.append(f"GitHub: {tailored.github_url}")
    lines.append("")

    if tailored.summary:
        lines.append("## Summary")
        lines.append(tailored.summary)
        lines.append("")

    if tailored.experiences:
        lines.append("## Experience")
        for exp in tailored.experiences:
            end = "Present" if exp.is_current else str(exp.end_date or "")
            lines.append(f"### {exp.title} — {exp.company}")
            lines.append(f"*{exp.start_date} to {end}*")
            for bullet in exp.bullets:
                lines.append(f"- {bullet}")
            lines.append("")

    if tailored.education:
        lines.append("## Education")
        for edu in tailored.education:
            field = f" in {edu.field_of_study}" if edu.field_of_study else ""
            lines.append(f"### {edu.degree}{field} — {edu.institution}")
            if edu.start_date or edu.end_date:
                lines.append(f"*{edu.start_date} to {edu.end_date or ''}*")
            lines.append("")

    if tailored.skills:
        lines.append("## Skills")
        for group in tailored.skills:
            lines.append(f"**{group.category}:** {', '.join(group.skills)}")
        lines.append("")

    if tailored.projects:
        lines.append("## Projects")
        for proj in tailored.projects:
            lines.append(f"### {proj.name}")
            if proj.description:
                lines.append(proj.description)
            if proj.technologies:
                lines.append(f"*Technologies: {', '.join(proj.technologies)}*")
            lines.append("")

    if tailored.certifications:
        lines.append("## Certifications")
        for cert in tailored.certifications:
            lines.append(f"- {cert}")
        lines.append("")

    return "\n".join(lines)


def to_docx(tailored: TailoredResume) -> bytes:
    """Render to DOCX bytes — ATS-safe: single column, standard headings, no tables."""
    from io import BytesIO

    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)

    doc.add_heading(tailored.full_name, level=0)

    contact_parts: list[str] = []
    if tailored.email:
        contact_parts.append(tailored.email)
    if tailored.phone:
        contact_parts.append(tailored.phone)
    if tailored.location:
        contact_parts.append(tailored.location)
    if contact_parts:
        doc.add_paragraph(" | ".join(contact_parts))

    if tailored.summary:
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(tailored.summary)

    if tailored.experiences:
        doc.add_heading("Experience", level=1)
        for exp in tailored.experiences:
            end = "Present" if exp.is_current else str(exp.end_date or "")
            doc.add_heading(f"{exp.title} — {exp.company}", level=2)
            doc.add_paragraph(f"{exp.start_date} to {end}")
            for bullet in exp.bullets:
                doc.add_paragraph(bullet, style="List Bullet")

    if tailored.education:
        doc.add_heading("Education", level=1)
        for edu in tailored.education:
            field = f" in {edu.field_of_study}" if edu.field_of_study else ""
            doc.add_heading(f"{edu.degree}{field} — {edu.institution}", level=2)

    if tailored.skills:
        doc.add_heading("Skills", level=1)
        for group in tailored.skills:
            doc.add_paragraph(f"{group.category}: {', '.join(group.skills)}")

    if tailored.projects:
        doc.add_heading("Projects", level=1)
        for proj in tailored.projects:
            doc.add_heading(proj.name, level=2)
            if proj.description:
                doc.add_paragraph(proj.description)

    if tailored.certifications:
        doc.add_heading("Certifications", level=1)
        for cert in tailored.certifications:
            doc.add_paragraph(cert, style="List Bullet")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_html(tailored: TailoredResume) -> str:
    """Render to ATS-plain HTML for PDF generation."""
    md = to_markdown(tailored)
    lines = md.split("\n")
    html_parts = ['<!DOCTYPE html><html><head><meta charset="utf-8">']
    html_parts.append("<style>")
    html_parts.append(
        "body{font-family:Arial,sans-serif;max-width:800px;margin:40px auto;"
        "padding:0 20px;font-size:11pt;line-height:1.4}"
    )
    html_parts.append("h1{font-size:18pt;margin-bottom:4px}")
    html_parts.append("h2{font-size:14pt;border-bottom:1px solid #ccc;padding-bottom:4px}")
    html_parts.append("h3{font-size:12pt;margin-bottom:2px}")
    html_parts.append("ul{margin-top:4px}")
    html_parts.append("</style></head><body>")

    for line in lines:
        if line.startswith("### "):
            html_parts.append(f"<h3>{_esc(line[4:])}</h3>")
        elif line.startswith("## "):
            html_parts.append(f"<h2>{_esc(line[3:])}</h2>")
        elif line.startswith("# "):
            html_parts.append(f"<h1>{_esc(line[2:])}</h1>")
        elif line.startswith("- "):
            html_parts.append(f"<li>{_esc(line[2:])}</li>")
        elif line.startswith("*") and line.endswith("*"):
            html_parts.append(f"<p><em>{_esc(line[1:-1])}</em></p>")
        elif line.startswith("**"):
            html_parts.append(f"<p><strong>{_esc(line.replace('**', ''))}</strong></p>")
        elif line.strip():
            html_parts.append(f"<p>{_esc(line)}</p>")

    html_parts.append("</body></html>")
    return "\n".join(html_parts)


def _esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


async def to_pdf(html_content: str) -> bytes:
    """Render HTML to PDF bytes via Playwright."""
    from playwright.async_api import async_playwright

    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page()
        await page.set_content(html_content, wait_until="networkidle")
        pdf_bytes = await page.pdf(
            format="Letter",
            margin={"top": "0.5in", "bottom": "0.5in", "left": "0.5in", "right": "0.5in"},
        )
        await browser.close()
    return pdf_bytes
