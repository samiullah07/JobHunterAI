"""Cover letter rendering — Markdown, DOCX, HTML, PDF outputs."""

from __future__ import annotations

from jobhunter.adapters.rendering.resume_renderer import to_pdf
from jobhunter.domain.cover_letter import CoverLetter


def to_markdown(letter: CoverLetter) -> str:
    lines: list[str] = []
    lines.append(letter.salutation)
    lines.append("")
    lines.append(letter.opening)
    lines.append("")
    for para in letter.body_paragraphs:
        lines.append(para)
        lines.append("")
    lines.append(letter.closing)
    lines.append("")
    lines.append(letter.signature)
    return "\n".join(lines)


def to_docx(letter: CoverLetter) -> bytes:
    """Render to DOCX bytes — ATS-safe: standard business-letter layout, no tables."""
    from io import BytesIO

    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)

    doc.add_paragraph(letter.salutation)
    doc.add_paragraph(letter.opening)
    for para in letter.body_paragraphs:
        doc.add_paragraph(para)
    doc.add_paragraph(letter.closing)
    doc.add_paragraph(letter.signature)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_html(letter: CoverLetter) -> str:
    """Render to HTML for PDF generation."""
    parts = ['<!DOCTYPE html><html><head><meta charset="utf-8">']
    parts.append("<style>")
    parts.append(
        "body{font-family:Arial,sans-serif;max-width:700px;margin:40px auto;"
        "padding:0 20px;font-size:11pt;line-height:1.6}"
    )
    parts.append("p{margin-bottom:12px}")
    parts.append("</style></head><body>")

    parts.append(f"<p>{_esc(letter.salutation)}</p>")
    parts.append(f"<p>{_esc(letter.opening)}</p>")
    for para in letter.body_paragraphs:
        parts.append(f"<p>{_esc(para)}</p>")
    parts.append(f"<p>{_esc(letter.closing)}</p>")
    parts.append(f"<p>{_esc(letter.signature)}</p>")

    parts.append("</body></html>")
    return "\n".join(parts)


def _esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


__all__ = ["to_docx", "to_html", "to_markdown", "to_pdf"]
